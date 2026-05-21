#!/usr/bin/env python3
"""
Build the staging-audit XLSX from a JSON data file and a directory of composed images.

Usage:
  python create_audit.py --data audit_data.json --images-dir ./audit_images/ --output staging-audit-2026-05-21.xlsx

audit_data.json schema:
[
  {
    "frame": "Login Screen",
    "issue": "Submit button background is #E63946 in Figma, #FF0000 on staging",
    "priority": "Major",
    "image_file": "row_1_image.png"
  },
  ...
]
"""
import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import (
    Alignment, Font, PatternFill, Border, Side
)
from openpyxl.utils import get_column_letter

HEADERS = ["Sr No", "Occurred on", "Image reference", "Issue identified",
           "Priority", "Dev comments", "Designer sign-off"]

COL_WIDTHS = [8, 25, 62, 50, 15, 35, 25]
ROW_HEIGHT = 120

HEADER_BG   = "1A1A2E"
HEADER_FG   = "FFFFFF"
ROW_A_BG    = "F5F5F5"
ROW_B_BG    = "FFFFFF"
EMPTY_BG    = "EEEEEE"

PRIORITY_COLORS = {
    "Critical": ("FF4444", "FFFFFF"),
    "Major":    ("FF9900", "000000"),
    "Minor":    ("FFD700", "000000"),
}

thin = Side(style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def header_style():
    return Font(bold=True, color=HEADER_FG, name="Arial", size=10)


def cell_font():
    return Font(name="Arial", size=9)


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", required=True)
    ap.add_argument("--images-dir", required=True)
    ap.add_argument("--output", required=True)
    args = ap.parse_args()

    with open(args.data) as f:
        rows = json.load(f)

    images_dir = Path(args.images_dir)
    wb = Workbook()
    ws = wb.active
    ws.title = "Staging Audit"

    # Header row
    for col, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_style()
        cell.fill = fill(HEADER_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.row_dimensions[1].height = 30

    # Data rows
    for idx, row in enumerate(rows, start=1):
        excel_row = idx + 1
        row_bg = ROW_A_BG if idx % 2 == 0 else ROW_B_BG

        values = [
            idx,
            row.get("frame", ""),
            "",  # image reference — handled separately
            row.get("issue", ""),
            row.get("priority", ""),
            "",
            "",
        ]

        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=excel_row, column=col, value=val)
            cell.font = cell_font()
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True,
                                       horizontal="center" if col in (1, 5) else "left")

            # Row background
            if col in (6, 7):
                cell.fill = fill(EMPTY_BG)
            else:
                cell.fill = fill(row_bg)

            # Priority color coding
            if col == 5:
                priority = str(val)
                if priority in PRIORITY_COLORS:
                    bg, fg = PRIORITY_COLORS[priority]
                    cell.fill = fill(bg)
                    cell.font = Font(name="Arial", size=9, bold=True, color=fg)

        ws.row_dimensions[excel_row].height = ROW_HEIGHT

        # Embed image
        img_file = row.get("image_file")
        if img_file:
            img_path = images_dir / img_file
            if img_path.exists():
                xl_img = XLImage(str(img_path))
                xl_img.width = 460
                xl_img.height = 110
                img_cell = f"C{excel_row}"
                ws.add_image(xl_img, img_cell)

    output_path = Path(args.output)
    wb.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    main()
